"""
sincronizar.py — Vuelca los movimientos directamente en tu fichero de contabilidad.

Escribe SOLO en la hoja de datos que le indiques. El resto del libro (fórmulas,
formato condicional, celdas combinadas, colores, anchos, gráficos, imágenes,
validación de datos) se conserva: está comprobado.

Lo que NO sobrevive a una reescritura son las tablas dinámicas y las macros.
Por eso, antes de tocar nada, el módulo inspecciona el libro y se niega a
escribir si encuentra alguna de las dos.

Se configura en sincronizar.json. Si ese fichero no existe, no se hace nada.
"""

from __future__ import annotations

import datetime as dt
import json
import os
import shutil
import zipfile
from pathlib import Path

import pandas as pd

# Dónde van las copias de seguridad. process.py la fija a datos/copias/ al
# arrancar; el valor por defecto solo sirve si alguien importa este módulo
# suelto para probar algo.
CARPETA_COPIAS = "copias"


def fijar_carpeta_copias(carpeta) -> None:
    global CARPETA_COPIAS
    CARPETA_COPIAS = str(carpeta)


# =====================================================================
# CONFIGURACIÓN
# =====================================================================

class Config:
    def __init__(self, datos: dict, raiz=None):
        self.archivo = str(datos.get("archivo", "")).strip()
        # Quien escriba «contabilidad.xlsx» en el JSON se refiere al de la
        # carpeta del programa, no al del directorio desde el que se lance la
        # consola. Las rutas absolutas se respetan tal cual.
        self.raiz = os.path.abspath(raiz) if raiz else os.path.abspath(".")
        self.hoja = datos.get("hoja", "MOVIMIENTOS")
        self.fila_inicial = int(datos.get("fila_inicial", 1))
        self.columna_inicial = int(datos.get("columna_inicial", 1))
        self.incluir_excluidos = bool(datos.get("incluir_excluidos", False))
        self.copias_de_seguridad = int(datos.get("copias_de_seguridad", 10))
        # «tabla» (lo de siempre): la hoja es de la herramienta y se reescribe
        # entera. «añadir»: la hoja es TUYA; solo se añaden debajo los
        # movimientos que aún no están, y lo que ya hay no se toca nunca.
        self.modo = str(datos.get("modo", "tabla")).strip().lower() or "tabla"
        # {"Tu cabecera": "campo"}, en el orden en que van. None = las siete
        # de siempre con su nombre.
        self.columnas = datos.get("columnas") or None

    @property
    def activa(self) -> bool:
        return bool(self.archivo)

    @property
    def ruta(self) -> str:
        """El fichero de destino, ya resuelto contra la raíz del programa."""
        if not self.archivo:
            return ""
        if os.path.isabs(self.archivo):
            return self.archivo
        return os.path.join(self.raiz, self.archivo)

    @classmethod
    def desde_json(cls, ruta, raiz=None):
        if not os.path.exists(ruta):
            return cls({}, raiz)
        from reglas import leer_json
        datos = {k: v for k, v in leer_json(ruta).items() if not k.startswith("_")}
        return cls(datos, raiz)


# =====================================================================
# FICHEROS ABIERTOS
# =====================================================================
# Viven aquí y no en process.py porque los usan los dos: el histórico y el
# fichero de contabilidad del usuario. Con este último, Windows solo falla
# al GUARDAR, y Mac y Linux ni eso: se escribía encima de un fichero abierto
# en LibreOffice, que al guardar después pisaba lo volcado sin avisar.

def _ficheros_de_bloqueo(ruta):
    """Los que deja al lado quien tiene el fichero abierto. Excel: «~$» más
    el nombre (con los dos primeros caracteres recortados en nombres
    largos, según versión). LibreOffice y OnlyOffice: «.~lock.nombre#»."""
    nombre = ruta.name
    return [ruta.with_name(n) for n in
            (f"~${nombre}", f"~${nombre[2:]}", f".~lock.{nombre}#")]


def esta_abierto(ruta):
    """
    None si se puede escribir sin miedo. Si no, cómo se ha sabido:
    "sistema" (Windows lo tiene bloqueado: seguro que está abierto) o el
    fichero de bloqueo encontrado (casi seguro, pero puede ser un resto de
    un programa que se cerró en falso y no lo borró).
    """
    ruta = Path(ruta)
    if not ruta.exists():
        return None
    if os.name == "nt":
        try:
            with open(ruta, "r+b"):
                pass
        except PermissionError:
            return "sistema"
        except OSError:
            pass
    for bloqueo in _ficheros_de_bloqueo(ruta):
        if bloqueo.exists():
            return bloqueo
    return None


# =====================================================================
# COMPROBACIONES PREVIAS
# =====================================================================

def revisar_libro(ruta: str) -> tuple[list[str], list[str]]:
    """
    Mira dentro del .xlsx antes de tocarlo.
    Devuelve (motivos_para_no_escribir, avisos_informativos).
    """
    if not os.path.exists(ruta):
        return ([f"No encuentro el fichero de contabilidad: {ruta}"], [])

    ext = os.path.splitext(ruta)[1].lower()
    if ext == ".xlsm":
        return (["El fichero es .xlsm (con macros). Las macros se perderían al "
                 "reescribirlo. Guárdalo como .xlsx si no las usas."], [])
    if ext != ".xlsx":
        return ([f"Solo sé escribir en ficheros .xlsx, y este es {ext}. "
                 f"Ábrelo y guárdalo como .xlsx."], [])

    try:
        nombres = zipfile.ZipFile(ruta).namelist()
    except zipfile.BadZipFile:
        return ([f"'{ruta}' no parece un .xlsx válido."], [])

    bloqueos, avisos = [], []
    if any("pivotTable" in n or "pivotCache" in n for n in nombres):
        bloqueos.append("El libro tiene TABLAS DINÁMICAS: se perderían al reescribirlo.")
    if any("vbaProject" in n for n in nombres):
        bloqueos.append("El libro tiene MACROS: se perderían al reescribirlo.")

    if any(n.startswith("xl/charts/") for n in nombres):
        avisos.append("tiene gráficos (se conservan, pero repásalos la primera vez)")
    if any(n.startswith("xl/media/") for n in nombres):
        avisos.append("tiene imágenes (se conservan, pero repásalas la primera vez)")

    return bloqueos, avisos


def revisar_hoja(ws) -> list[str]:
    """
    La hoja de destino tiene que ser de datos, no de cálculo.

    Este es el seguro más importante: si por una errata apuntas a la hoja de
    totales, aquí se detecta que hay fórmulas y no se borra nada.

    Se mira la hoja ENTERA, no solo el bloque que se va a volcar: una fórmula
    en cualquier parte de la hoja significa que no es una hoja de datos, y a la
    mínima duda es mejor no escribir. Por eso no recibe la esquina del bloque.
    """
    formulas = []
    for fila in ws.iter_rows():
        for celda in fila:
            if celda.data_type == "f":
                formulas.append(celda.coordinate)
                if len(formulas) >= 5:
                    break
        if len(formulas) >= 5:
            break

    if formulas:
        return [f"La hoja «{ws.title}» contiene fórmulas ({', '.join(formulas)}...). "
                f"No la sobreescribo: la hoja de destino debe ser solo datos. "
                f"Comprueba el nombre de la hoja en sincronizar.json."]
    return []


MODOS = ("tabla", "añadir")
# lo mínimo para reconocer un movimiento en la hoja: sin esto no se pueden
# recolocar las notas (tabla) ni saber qué falta por añadir (añadir)
CAMPOS_OBLIGATORIOS = ("fecha", "descripcion", "importe")


def columnas_a_volcar(cfg: Config, disponibles: list[str]) -> tuple[list, list]:
    """(campos, cabeceras): qué columnas del resultado se escriben y con qué
    texto de cabecera. Un error de configuración para en seco, con el
    fichero sin tocar."""
    if cfg.modo not in MODOS:
        raise RuntimeError(f"En sincronizar.json, «modo» tiene que ser "
                           f"\"tabla\" o \"añadir\", no «{cfg.modo}».")
    if cfg.columnas is None:
        return list(disponibles), list(disponibles)
    if not isinstance(cfg.columnas, dict) or not cfg.columnas:
        raise RuntimeError("En sincronizar.json, «columnas» tiene que ser "
                           "{\"Tu cabecera\": \"campo\", ...}.")
    cabeceras = [str(k) for k in cfg.columnas]
    campos = [str(v).strip().lower() for v in cfg.columnas.values()]
    malos = [c for c in campos if c not in disponibles]
    if malos:
        raise RuntimeError(f"En sincronizar.json, «columnas» usa {', '.join(malos)}, "
                           f"que no es ninguno de los campos que vuelco: "
                           f"{', '.join(disponibles)}.")
    if len(set(campos)) < len(campos):
        raise RuntimeError("En sincronizar.json, «columnas» repite un campo.")
    faltan = [c for c in CAMPOS_OBLIGATORIOS if c not in campos]
    if faltan:
        raise RuntimeError(f"En sincronizar.json, a «columnas» le falta "
                           f"{', '.join(faltan)}: sin fecha, concepto e importe "
                           f"no puedo reconocer tus movimientos en la hoja.")
    return campos, cabeceras


def _igual(a, b) -> bool:
    from reglas import normalizar
    return normalizar(a) == normalizar(b)


def revisar_esquina(ws, f0, c0, columnas) -> list[str]:
    """
    El bloque se vacía y se reescribe ENTERO en cada ejecución: es una tabla
    de la herramienta, no se añade debajo de lo que haya. Si en la esquina
    configurada ya hay datos del usuario (su hoja de siempre, con junio
    metido a mano y sus propias cabeceras), escribir se los llevaba por
    delante, y las fórmulas que tiraban de esas columnas pasaban a sumar
    otra cosa sin avisar.

    Se reconoce una tabla nuestra por su cabecera: cada celda escrita en la
    fila de la esquina tiene que ser el nombre de la columna que va ahí. Una
    fila de cabecera vacía con datos debajo tampoco es nuestra.
    """
    ancho = len(columnas)
    cabecera = [ws.cell(f0, c0 + j).value for j in range(ancho)]
    ajenas = [(j, v) for j, v in enumerate(cabecera)
              if v is not None and str(v).strip() and not _igual(v, columnas[j])]
    hay_debajo = any(celda.value is not None
                     for fila in ws.iter_rows(min_row=f0 + 1, min_col=c0,
                                              max_col=c0 + ancho - 1)
                     for celda in fila)
    vacia = all(v is None or not str(v).strip() for v in cabecera)
    if not ajenas and not (vacia and hay_debajo):
        return []

    from openpyxl.utils import get_column_letter
    esquina = f"{get_column_letter(c0)}{f0}"
    if ajenas:
        vistas = ", ".join(f"«{v}»" for _, v in ajenas[:4])
        que = f"una cabecera que no es la mía ({vistas})"
    else:
        que = "datos sin la cabecera de la herramienta"
    return [f"La hoja «{ws.title}» tiene {que} en la esquina {esquina}, donde "
            f"empiezo a escribir. Vuelco la tabla ENTERA en cada ejecución (no "
            f"añado debajo): si escribiera, borraría lo tuyo.",
            f"Usa una hoja vacía solo para la herramienta y haz que tus "
            f"fórmulas tiren de ella, o mueve la esquina (fila_inicial, "
            f"columna_inicial) a un sitio libre. Las columnas que escribo son: "
            f"{', '.join(columnas)}.",
            "Si es tu hoja de siempre y quieres que los movimientos nuevos se "
            "añadan debajo, pon \"modo\": \"añadir\" y en \"columnas\" tus "
            "cabeceras (la guía lo explica)."]


# =====================================================================
# COPIA DE SEGURIDAD
# =====================================================================

def ultima_copia(ruta) -> str | None:
    """La copia más reciente de ese fichero, si hay alguna."""
    if not os.path.isdir(CARPETA_COPIAS):
        return None
    base = os.path.splitext(os.path.basename(str(ruta)))[0]
    previas = sorted(f for f in os.listdir(CARPETA_COPIAS)
                     if f.startswith(base + "_") and f.endswith(".xlsx"))
    return os.path.join(CARPETA_COPIAS, previas[-1]) if previas else None


def copia_de_seguridad(ruta: str, maximo: int) -> str:
    os.makedirs(CARPETA_COPIAS, exist_ok=True)
    base = os.path.splitext(os.path.basename(ruta))[0]
    sello = dt.datetime.now().strftime("%Y%m%d_%H%M%S")
    destino = os.path.join(CARPETA_COPIAS, f"{base}_{sello}.xlsx")
    shutil.copy2(ruta, destino)

    if maximo > 0:
        previas = sorted(f for f in os.listdir(CARPETA_COPIAS)
                         if f.startswith(base + "_") and f.endswith(".xlsx"))
        for vieja in previas[:-maximo]:
            os.remove(os.path.join(CARPETA_COPIAS, vieja))

    return destino


# =====================================================================
# ESCRITURA
# =====================================================================

def _limpiar_bloque(ws, fila_inicial, columna_inicial, ancho):
    """Vacía lo que hubiera antes, sin tocar los estilos ni el resto de la hoja."""
    for fila in ws.iter_rows(min_row=fila_inicial,
                             min_col=columna_inicial,
                             max_col=columna_inicial + ancho - 1):
        for celda in fila:
            celda.value = None


def _hay_datos_fuera_del_bloque(ws, columna_inicial, ancho, desde_fila) -> bool:
    """
    ¿Hay algo escrito a los lados de las columnas que volcamos, de la fila
    `desde_fila` hacia abajo?

    Importa porque el recorte de filas sobrantes borra la FILA ENTERA, no solo
    el bloque. Si el usuario tiene una columna de notas al lado de los datos,
    borrar la fila se la lleva por delante.
    """
    primera = columna_inicial
    ultima = columna_inicial + ancho - 1
    for fila in ws.iter_rows(min_row=desde_fila):
        for celda in fila:
            if celda.value is None:
                continue
            if celda.column < primera or celda.column > ultima:
                return True
    return False


# --- notas del usuario al lado de los movimientos ---
# El bloque se vacía y se reescribe en orden de fecha, pero lo que el usuario
# escribe a los lados («regalo mamá» junto a un cargo) se quedaba en su número
# de fila. Bastaba con que entrara un movimiento más antiguo para que todo lo
# de debajo bajara una fila y cada nota acabara junto al movimiento de al
# lado, sin ningún aviso. Ahora cada nota se recuerda por el movimiento al que
# acompaña y se vuelve a poner a su lado.

def _clave_fila(fecha, descripcion, importe):
    """(fecha, concepto, importe) de una fila, o None si la fila no parece un
    movimiento (sin fecha o sin importe): lo que el usuario tenga junto a
    otras cosas no se toca.

    Admite lo que alguien teclea a mano en su hoja («02/06/2026», «-64,35»)
    y compara el concepto sin mayúsculas ni acentos: en modo «añadir» las
    filas de junio metidas a mano tienen que reconocerse como las mismas que
    trae el banco, o se añadirían otra vez."""
    from bank_io import parsear_fecha, parsear_importe
    from reglas import normalizar
    if not hasattr(fecha, "strftime"):
        fecha = parsear_fecha(fecha) if isinstance(fecha, str) else None
        if fecha is None or pd.isna(fecha):
            return None
    imp = parsear_importe(importe) if isinstance(importe, str) else importe
    try:
        imp = round(float(imp), 2)
    except (TypeError, ValueError):
        return None
    if pd.isna(imp):
        return None
    return fecha.strftime("%Y-%m-%d"), normalizar(descripcion or ""), imp


def _con_repeticion(claves):
    """Dos movimientos idénticos el mismo día son dos: se numeran, igual que
    en el histórico, para que cada nota vuelva a SU movimiento."""
    vistos, salida = {}, []
    for k in claves:
        if k is None:
            salida.append(None)
            continue
        n = vistos.get(k, 0)
        vistos[k] = n + 1
        salida.append((k, n))
    return salida


def _leer_notas(ws, f0, c0, ancho, cols):
    """{movimiento: (fila, {columna: valor})} de lo escrito a los lados."""
    if not all(c in cols for c in ("fecha", "descripcion", "importe")):
        return {}
    cf, cd, ci = (c0 + cols.index(c) for c in ("fecha", "descripcion", "importe"))
    filas = list(range(f0 + 1, ws.max_row + 1))
    claves = _con_repeticion([_clave_fila(ws.cell(r, cf).value, ws.cell(r, cd).value,
                                          ws.cell(r, ci).value) for r in filas])
    notas = {}
    for r, k in zip(filas, claves):
        if k is None:
            continue
        lados = {c.column: c.value for c in ws[r]
                 if c.value is not None and not c0 <= c.column < c0 + ancho}
        if lados:
            notas[k] = (r, lados)
    return notas


def _recolocar_notas(ws, notas, df, f0):
    """Pone cada nota junto a su movimiento. Nunca pisa otra: si el sitio
    está ocupado, la deja donde estaba. Devuelve (movidas, sin_su_movimiento,
    sin_sitio)."""
    claves = _con_repeticion([_clave_fila(r["fecha"], r["descripcion"], r["importe"])
                              for _, r in df.iterrows()])
    destino = {k: f0 + 1 + i for i, k in enumerate(claves) if k is not None}
    mover = {k: v for k, v in notas.items() if k in destino and destino[k] != v[0]}
    huerfanas = sum(1 for k in notas if k not in destino)

    for r, lados in mover.values():
        for c in lados:
            ws.cell(r, c).value = None
    sin_sitio = 0
    for k, (r, lados) in mover.items():
        fila = destino[k]
        if any(ws.cell(fila, c).value is not None for c in lados):
            sin_sitio += 1
            if all(ws.cell(r, c).value is None for c in lados):
                for c, v in lados.items():
                    ws.cell(r, c).value = v
            continue
        for c, v in lados.items():
            ws.cell(fila, c).value = v
    return len(mover) - sin_sitio, huerfanas, sin_sitio


def _contenido(ws) -> dict:
    """Lo que hay escrito en la hoja, celda a celda: para saber si volcar
    cambia algo antes de guardar (y de gastar una copia de seguridad)."""
    return {(c.row, c.column): c.value for fila in ws.iter_rows()
            for c in fila if c.value is not None}


def escribir(cfg: Config, df: pd.DataFrame) -> tuple[str | None, int]:
    """Vuelca df en la hoja indicada. Devuelve (ruta_de_la_copia, filas
    escritas); la copia es None si no había nada que cambiar y el fichero
    no se ha tocado."""
    from openpyxl import load_workbook

    bloqueos, avisos = revisar_libro(cfg.ruta)
    if bloqueos:
        raise RuntimeError("No sincronizo con tu fichero de contabilidad:\n   · "
                           + "\n   · ".join(bloqueos))

    abierto = esta_abierto(cfg.ruta)
    if abierto:
        como = ("" if abierto == "sistema" else
                f"\n   (lo sé por {abierto.name}; si NO lo tienes abierto, es un "
                f"resto de un cierre en falso: bórralo)")
        raise RuntimeError(f"'{cfg.archivo}' está abierto en otro programa. "
                           f"Guárdalo, ciérralo y vuelve a ejecutar: si escribo "
                           f"ahora, al guardar tú después se perdería lo "
                           f"volcado.{como}")

    campos, cabeceras = columnas_a_volcar(cfg, [str(c) for c in df.columns])
    df = df[campos]

    for a in avisos:
        print(f"   ℹ️  El libro {a}")

    try:
        wb = load_workbook(cfg.ruta)
    except PermissionError:
        raise RuntimeError(f"'{cfg.archivo}' está abierto en otro programa. "
                           f"Ciérralo y vuelve a ejecutar.")

    if cfg.hoja not in wb.sheetnames:
        wb.close()
        raise RuntimeError(f"El libro no tiene ninguna hoja llamada «{cfg.hoja}». "
                           f"Hojas disponibles: {', '.join(wb.sheetnames)}.")

    ws = wb[cfg.hoja]
    revisar = revisar_cabecera_propia if cfg.modo == "añadir" else revisar_esquina
    problemas = revisar_hoja(ws) or revisar(
        ws, cfg.fila_inicial, cfg.columna_inicial, cabeceras)
    if problemas:
        wb.close()
        raise RuntimeError("No sincronizo:\n   · " + "\n   · ".join(problemas))

    antes = _contenido(ws)
    if cfg.modo == "añadir":
        nuevos, anteriores, cats_nuevas = _anadir(
            ws, cfg.fila_inicial, cfg.columna_inicial, df, cabeceras)
        if _contenido(ws) == antes:
            wb.close()
            return None, 0
        copia = copia_de_seguridad(cfg.ruta, cfg.copias_de_seguridad)
        if cats_nuevas:
            print(f"   ℹ️  Categorías que tu hoja no tenía hasta ahora: "
                  + ", ".join(f"{c} ({n})" for c, n in cats_nuevas)
                  + ".\n      Si tus totales van por categoría, añádelas o no "
                    "sumarán en ninguno.")
        if anteriores:
            print(f"   ℹ️  {anteriores} de los añadidos son anteriores al último "
                  f"que ya tenías: van al final. Ordena la hoja por fecha si "
                  f"quieres verlos en su sitio.")
        _guardar(wb, cfg, copia)
        return copia, nuevos

    notas = _leer_notas(ws, cfg.fila_inicial, cfg.columna_inicial,
                        len(df.columns), list(df.columns))
    _limpiar_bloque(ws, cfg.fila_inicial, cfg.columna_inicial, len(df.columns))

    f0, c0 = cfg.fila_inicial, cfg.columna_inicial
    for j, cabecera in enumerate(cabeceras):
        ws.cell(row=f0, column=c0 + j, value=cabecera)

    for i, (_, registro) in enumerate(df.iterrows(), start=1):
        _escribir_fila(ws, f0 + i, c0, registro)

    movidas = huerfanas = sin_sitio = 0
    if notas:
        movidas, huerfanas, sin_sitio = _recolocar_notas(ws, notas, df, f0)

    # Si esta vez hay menos filas que la anterior, no basta con vaciarlas: se
    # eliminan, o la hoja arrastra para siempre el rango usado más grande.
    #
    # Pero delete_rows borra la FILA ENTERA. Si el usuario tiene datos propios
    # a la derecha (o a la izquierda) del bloque volcado, se los llevaría por
    # delante. En ese caso se prefiere dejar las filas: ya han quedado vacías
    # en las columnas que nos tocan, y lo único que se pierde es el recorte del
    # rango usado, que es cosmético. Perder las notas del usuario no lo es.
    ultima = f0 + len(df)
    filas_conservadas = False
    if ws.max_row > ultima:
        if _hay_datos_fuera_del_bloque(ws, c0, len(df.columns), ultima + 1):
            filas_conservadas = True
        else:
            ws.delete_rows(ultima + 1, ws.max_row - ultima)

    # Sin nada nuevo, ni se guarda ni se copia. Antes cada ejecución gastaba
    # una copia de seguridad aunque no cambiara nada, y a las diez la única
    # anterior a la primera sincronización (la que tiene tus datos de antes)
    # se borraba sola.
    if _contenido(ws) == antes:
        wb.close()
        return None, len(df)

    copia = copia_de_seguridad(cfg.ruta, cfg.copias_de_seguridad)
    if movidas:
        print(f"   ℹ️  {movidas} notas tuyas recolocadas junto a su movimiento "
              f"(han entrado movimientos por medio).")
    if huerfanas or sin_sitio:
        print(f"   ⚠️  {huerfanas + sin_sitio} notas tuyas no he podido "
              f"ponerlas junto a su movimiento (ya no se vuelca, o el sitio\n"
              f"      estaba ocupado): se han quedado donde estaban. Están "
              f"como antes en la copia {copia}.")
    if filas_conservadas:
        print(f"   ℹ️  Por debajo de la fila {ultima} hay datos tuyos fuera "
              f"de las columnas que vuelco.")
        print(f"      He vaciado esas filas en el bloque volcado, pero no las "
              f"he borrado, para no llevarme lo demás.")

    _guardar(wb, cfg, copia)
    return copia, len(df)


def _guardar(wb, cfg, copia):
    try:
        wb.save(cfg.ruta)
    except PermissionError:
        raise RuntimeError(f"No he podido guardar '{cfg.archivo}': está abierto en "
                           f"otro programa. Tu copia de seguridad está en {copia}.")
    finally:
        wb.close()


def _escribir_fila(ws, fila, c0, registro):
    for j, (nombre, valor) in enumerate(registro.items()):
        celda = ws.cell(row=fila, column=c0 + j)
        if isinstance(valor, pd.Timestamp):
            celda.value = valor.to_pydatetime()
            celda.number_format = "DD/MM/YYYY"
        elif pd.isna(valor):
            celda.value = None
        elif isinstance(valor, (int, float)) and not isinstance(valor, bool):
            celda.value = float(valor)
            if nombre == "importe":
                celda.number_format = '#,##0.00 €'
        else:
            texto = str(valor)
            celda.value = texto
            if texto.startswith("="):        # que no lo tome por fórmula
                celda.data_type = "s"


# --- modo «añadir»: la hoja es del usuario ---
# Para quien lleva su contabilidad en su propia hoja desde hace años (con
# meses metidos a mano y sus cabeceras) y solo quiere dejar de copiar y
# pegar. Nada de lo que ya hay se mueve, se reescribe ni se borra: se
# reconoce cada movimiento por fecha, concepto e importe, y solo se añaden
# debajo los que faltan. El precio: lo ya añadido no se reclasifica si luego
# cambias una regla (es tuyo, como lo que tecleas a mano).

def revisar_cabecera_propia(ws, f0, c0, cabeceras) -> list[str]:
    """La cabecera de la esquina tiene que ser la de «columnas», o estar
    vacía con la hoja vacía debajo (primera vez)."""
    from openpyxl.utils import get_column_letter
    ancho = len(cabeceras)
    vistas = [ws.cell(f0, c0 + j).value for j in range(ancho)]
    if all(v is None or not str(v).strip() for v in vistas):
        hay_debajo = any(c.value is not None
                         for fila in ws.iter_rows(min_row=f0 + 1, min_col=c0,
                                                  max_col=c0 + ancho - 1)
                         for c in fila)
        if not hay_debajo:
            return []
        return [f"La hoja «{ws.title}» tiene datos pero no cabecera en la fila "
                f"{f0}: pon ahí tus cabeceras ({', '.join(cabeceras)}) o "
                f"cambia fila_inicial a la fila donde las tengas."]
    distintas = [f"{get_column_letter(c0 + j)}{f0}: «{v or ''}» en vez de «{e}»"
                 for j, (v, e) in enumerate(zip(vistas, cabeceras))
                 if not _igual(v or "", e)]
    if not distintas:
        return []
    return [f"La cabecera de «{ws.title}» no coincide con «columnas» de "
            f"sincronizar.json ({'; '.join(distintas[:4])}). No añado nada: "
            f"escribiría cada dato en la columna que no es.",
            "Pon en «columnas» tus cabeceras tal cual, en su orden, y a qué "
            "campo corresponde cada una."]


def _anadir(ws, f0, c0, df, cabeceras) -> tuple[int, int, list]:
    """Añade debajo lo que no está. Devuelve (añadidos, cuántos de ellos son
    anteriores al último que ya había, categorías nuevas en la hoja)."""
    ancho = len(cabeceras)
    campos = list(df.columns)
    if all(ws.cell(f0, c0 + j).value in (None, "") for j in range(ancho)):
        for j, cabecera in enumerate(cabeceras):
            ws.cell(row=f0, column=c0 + j, value=cabecera)

    cf, cd, ci = (c0 + campos.index(c) for c in CAMPOS_OBLIGATORIOS)
    ultima, claves = f0, []
    for r in range(f0 + 1, ws.max_row + 1):
        if any(ws.cell(r, c0 + j).value not in (None, "") for j in range(ancho)):
            ultima = r
        claves.append(_clave_fila(ws.cell(r, cf).value, ws.cell(r, cd).value,
                                  ws.cell(r, ci).value))
    ya = {k for k in _con_repeticion(claves) if k is not None}
    mas_reciente = max((k[0][0] for k in ya), default="")

    # Las filas nuevas con el formato de la última que ya había (su formato
    # de fecha, sus euros): si no, la hoja quedaba con dos aspectos.
    formatos = ({j: ws.cell(ultima, c0 + j).number_format for j in range(ancho)}
                if ultima > f0 else {})
    # Qué categorías usaba ya la hoja: una nueva (Higiene, que la hoja nunca
    # tuvo) no la recoge ningún total que vaya por categoría, y se dice.
    col_cat = c0 + campos.index("categoria") if "categoria" in campos else None
    usadas = ({str(ws.cell(r, col_cat).value).strip() for r in range(f0 + 1, ultima + 1)
               if ws.cell(r, col_cat).value not in (None, "")}
              if col_cat and ultima > f0 else None)
    nuevas = {}

    df = df.sort_values("fecha", kind="stable")
    entrantes = _con_repeticion([_clave_fila(r["fecha"], r["descripcion"], r["importe"])
                                 for _, r in df.iterrows()])
    anadidos = anteriores = 0
    for (_, registro), k in zip(df.iterrows(), entrantes):
        if k is None or k in ya:
            continue
        ultima += 1
        _escribir_fila(ws, ultima, c0, registro)
        for j, formato in formatos.items():
            ws.cell(ultima, c0 + j).number_format = formato
        anadidos += 1
        anteriores += k[0][0] < mas_reciente
        cat = str(registro.get("categoria", "") or "").strip()
        if usadas is not None and cat and cat not in usadas:
            nuevas[cat] = nuevas.get(cat, 0) + 1
    return anadidos, anteriores, sorted(nuevas.items())
