"""
fixtures.py — Fabrica ficheros de banco falsos, uno por cada formato que los
bancos españoles llaman «.xls».

Todos los datos son inventados. Las fechas son de 2026 y los importes son
números redondos a propósito, para que los totales esperados de las pruebas se
puedan comprobar a mano sin calculadora.

Cada generador recibe la misma lista de movimientos:

    [("05/04/2026", "NOMINA EMPRESA SL", 2000.00), ...]

y la escribe en el formato que le toca. Así el mismo caso de prueba se puede
pasar por los cinco lectores sin reescribir los datos.
"""

from __future__ import annotations

import io
from pathlib import Path

# =====================================================================
# HTML renombrado a .xls  (el formato más común en la banca española)
# =====================================================================

def escribir_html(ruta: Path, movimientos, cabecera_saldo=True):
    """Tabla HTML con preámbulo de banco y, opcionalmente, columna 'Saldo'
    (que es la pista por la que detectar_tipo() decide que es una cuenta)."""
    filas = [
        "<tr><td>BANCO FICTICIO S.A.</td><td></td><td></td><td></td></tr>",
        "<tr><td>Extracto de cuenta</td><td></td><td></td><td></td></tr>",
        "<tr><td>Cuenta:</td><td>ES00 0000 0000 0000 0000 0000</td><td></td><td></td></tr>",
        "<tr><td></td><td></td><td></td><td></td></tr>",
    ]
    if cabecera_saldo:
        filas.append("<tr><td>Fecha operación</td><td>Concepto</td>"
                     "<td>Importe</td><td>Saldo</td></tr>")
    else:
        filas.append("<tr><td>Fecha operación</td><td>Concepto</td>"
                     "<td>Importe</td><td></td></tr>")

    saldo = 5000.0
    for fecha, concepto, importe in movimientos:
        saldo += importe
        filas.append(
            f"<tr><td>{fecha}</td><td>{concepto}</td>"
            f"<td>{_euros(importe)}</td><td>{_euros(saldo)}</td></tr>")

    html = ('<html><head><meta http-equiv="Content-Type" '
            'content="text/html; charset=windows-1252"></head><body>\n'
            "<table>\n" + "\n".join(filas) + "\n</table>\n</body></html>\n")
    ruta.write_bytes(html.encode("cp1252", errors="replace"))


# =====================================================================
# SpreadsheetML 2003 renombrado a .xls
# =====================================================================

def escribir_xml_ss(ruta: Path, movimientos, tarjeta=True):
    """XML de Office 2003. Con cabecera 'Importe de la operación' para que
    detectar_tipo() lo tome por un extracto de tarjeta."""
    col_importe = "Importe de la operación" if tarjeta else "Importe"
    filas = [
        _fila_xml(["TARJETA CREDITO ****1234"]),
        _fila_xml(["Periodo: abril 2026"]),
        _fila_xml([]),
        _fila_xml(["Fecha operación", "Comercio", col_importe]),
    ]
    for fecha, concepto, importe in movimientos:
        filas.append(_fila_xml([fecha, concepto, _euros(importe)]))

    xml = ('<?xml version="1.0"?>\n'
           '<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet"\n'
           ' xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet">\n'
           '<Worksheet ss:Name="Movimientos"><Table>\n'
           + "\n".join(filas) +
           "\n</Table></Worksheet>\n</Workbook>\n")
    ruta.write_text(xml, encoding="utf-8")


def _fila_xml(celdas):
    if not celdas:
        return "<Row></Row>"
    cs = "".join(f'<Cell><Data ss:Type="String">{_escapar(c)}</Data></Cell>'
                 for c in celdas)
    return f"<Row>{cs}</Row>"


def _escapar(t):
    return (str(t).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))


# =====================================================================
# CSV renombrado a .xls
# =====================================================================

def escribir_csv(ruta: Path, movimientos, sep=";", fechas_iso=False,
                 encoding="utf-8"):
    """CSV con preámbulo. Con fechas_iso=True escribe aaaa-mm-dd, que es el
    caso que invertía día y mes antes de detectar el formato ISO."""
    lineas = ["BANCO FICTICIO S.A.", "Extracto descargado", "",
              sep.join(["Fecha operación", "Concepto", "Importe"])]
    for fecha, concepto, importe in movimientos:
        if fechas_iso:
            d, m, a = fecha.split("/")
            fecha = f"{a}-{m}-{d}"
        lineas.append(sep.join([fecha, concepto, _euros(importe)]))
    ruta.write_text("\n".join(lineas) + "\n", encoding=encoding)


# =====================================================================
# XLSX renombrado a .xls
# =====================================================================

def escribir_xlsx(ruta: Path, movimientos):
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Movimientos"
    ws.append(["BANCO FICTICIO S.A."])
    ws.append(["Extracto de cuenta"])
    ws.append([])
    ws.append(["Fecha operación", "Concepto", "Importe"])
    for fecha, concepto, importe in movimientos:
        ws.append([fecha, concepto, importe])
    wb.save(ruta)
    wb.close()


# =====================================================================
# .xls BIFF «de verdad» (contenedor OLE2)
# =====================================================================

def escribir_biff_falso(ruta: Path):
    """
    No genera un BIFF real: escribe la firma OLE2 para que detectar_formato()
    lo mande por la ruta 'xls_biff', y esa ruta se prueba con el stub de xlrd
    (pruebas/xlrd_stub.py), que devuelve siempre la misma tabla.

    Es lo único que se puede hacer sin la librería xlrd instalada, y prueba
    justo lo que interesa: que bank_io elige el lector correcto y sabe traducir
    los tipos de celda de xlrd (fecha serial, número, vacío).
    """
    ruta.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 1024)


# =====================================================================
# Libro de contabilidad de destino, para probar sincronizar.py
# =====================================================================

def escribir_destino(ruta: Path, hoja="MOVIMIENTOS", notas_en_columna=None,
                     formula=False, filas_previas=8):
    """
    Imita el fichero de contabilidad del usuario.

    notas_en_columna: letra de una columna a la DERECHA del bloque volcado con
    datos propios. Es lo que comprueba que el recorte de filas sobrantes no se
    lleve por delante lo que el usuario tenga al lado.
    """
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = hoja
    ws["A1"] = "fecha"
    for i in range(2, filas_previas + 2):
        ws.cell(row=i, column=1, value=f"relleno {i}")
    if notas_en_columna:
        ws[f"{notas_en_columna}1"] = "mis notas"
        for i in range(2, filas_previas + 2):
            ws.cell(row=i, column=ws[f"{notas_en_columna}1"].column,
                    value=f"nota {i}")
    if formula:
        ws["A20"] = "=SUM(C2:C10)"
    wb.save(ruta)
    wb.close()


# =====================================================================

def _euros(v):
    """1234.5 -> '1.234,50' (formato español, como lo escupen los bancos)."""
    entero, dec = f"{abs(v):.2f}".split(".")
    grupos = []
    while len(entero) > 3:
        grupos.insert(0, entero[-3:])
        entero = entero[:-3]
    grupos.insert(0, entero)
    s = ".".join(grupos) + "," + dec
    return ("-" + s) if v < 0 else s
