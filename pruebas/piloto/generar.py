"""
generar.py — Prepara un piloto de usuarios simulados (ver INSTRUCCIONES.md).

Crea una carpeta aislada por perfil con la release descomprimida, la web
(guía y reglas) y los extractos INVENTADOS de ese perfil en descargas/, tal
como los bajaría del banco. Cada perfil tiene un «banco» distinto para forzar
lectores y cabeceras diferentes.

Todo es sintético: ningún dato sale de ajustes/ ni de datos/.

    app/.venv/bin/python pruebas/piloto/generar.py [--ronda N] DESTINO [ZIP] [WEB]

--ronda  qué perfiles generar (1 o 2; por defecto, la última)
DESTINO  carpeta donde crear los perfiles (mejor el scratchpad de la sesión)
ZIP      release a probar; por defecto, el eledger_v*.zip más nuevo de la raíz
WEB      copia local de la web; por defecto, ../eledger-web

Ojo, aprendido en la ronda 1: con UNA compra al mes en el mismo súper, el
mismo día, el informe de cargos que se repiten la toma por un recibo. Eso es
un artefacto de los datos, no del programa: en un extracto real hay varias
compras al mes en el mismo sitio. Por eso los perfiles de la ronda 2 usan
frecuentes() y recibo() en vez de un día fijo por comercio.
"""
import random
import shutil
import zipfile
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook

RAIZ = Path(__file__).resolve().parents[2]
random.seed(2026)


def es(v):
    """1234.5 -> '1.234,50' (con signo menos si hace falta)."""
    s = f"{abs(v):,.2f}".replace(",", "_").replace(".", ",").replace("_", ".")
    return ("-" if v < 0 else "") + s


def dmy(d):
    return d.strftime("%d/%m/%Y")


def con_saldo(movs, inicial):
    movs = sorted(movs, key=lambda m: m[0])
    saldo, out = inicial, []
    for d, c, i in movs:
        saldo = round(saldo + i, 2)
        out.append((d, c, i, saldo))
    return out


def mes(m, dia):
    return date(2026, m, dia)


# ---------------------------------------------------------------- formatos
def html_xls(ruta, filas, cab, titulo="Extracto de movimientos"):
    tr = [f"<tr><td>{titulo}</td></tr>", "<tr><td>Titular: ***</td></tr>", "<tr></tr>",
          "<tr>" + "".join(f"<th>{h}</th>" for h in cab) + "</tr>"]
    for f in filas:
        tr.append("<tr>" + "".join(f"<td>{x}</td>" for x in f) + "</tr>")
    html = ('<html><head><meta charset="windows-1252"></head><body><table>\n'
            + "\n".join(tr) + "\n</table></body></html>")
    ruta.write_bytes(html.encode("cp1252", errors="replace"))


def xml_xls(ruta, filas, cab):
    def fila(cs):
        return "<Row>" + "".join(
            f'<Cell><Data ss:Type="String">{c}</Data></Cell>' for c in cs) + "</Row>"
    rows = [fila(["Movimientos de tarjeta"]), fila([]), fila(cab)] + [fila(f) for f in filas]
    ruta.write_text('<?xml version="1.0"?>\n<Workbook xmlns="urn:schemas-microsoft-com:office:spreadsheet" '
                    'xmlns:ss="urn:schemas-microsoft-com:office:spreadsheet">'
                    '<Worksheet ss:Name="Hoja1"><Table>' + "".join(rows)
                    + "</Table></Worksheet></Workbook>", encoding="utf-8")


def csv(ruta, filas, cab, sep=";", encoding="cp1252", preambulo=("Consulta de movimientos", "")):
    lineas = list(preambulo) + [sep.join(cab)] + [sep.join(map(str, f)) for f in filas]
    ruta.write_bytes(("\r\n".join(lineas) + "\r\n").encode(encoding, errors="replace"))


def xlsx(ruta, filas, cab):
    wb = Workbook()
    ws = wb.active
    ws.append(["Movimientos tarjeta"])
    ws.append([])
    ws.append(cab)
    for f in filas:
        ws.append(list(f))
    wb.save(ruta)


# ---------------------------------------------------------------- perfiles
def lucia(d):
    """No técnica. Cuenta en HTML .xls, tarjeta en XML .xls, y la cuenta paga
    la tarjeta con un recibo que nadie ha excluido."""
    cuenta, tarjeta = [], []
    for m in (7, 8, 9):
        cuenta += [(mes(m, 1) - timedelta(days=1) if m > 7 else mes(7, 1), "NOMINA COLEGIO SAN JOSE", 1650.00),
                   (mes(m, 3), "RECIBO ALQUILER VIVIENDA", -620.00),
                   (mes(m, 5), "RECIBO IBERDROLA CLIENTES", -round(random.uniform(45, 80), 2)),
                   (mes(m, 8), "RECIBO CANAL DE ISABEL II", -round(random.uniform(20, 35), 2)),
                   (mes(m, 10), "RECIBO MOVISTAR", -42.90),
                   (mes(m, 14), "COMPRA MERCADONA", -round(random.uniform(60, 110), 2)),
                   (mes(m, 22), "RETIRADA CAJERO", -50.00)]
        suma = 0
        for dia, com, lo, hi in ((4, "CARREFOUR MARKET", 30, 90), (11, "FARMACIA CENTRAL", 8, 25),
                                 (17, "ZARA", 25, 80), (21, "PELUQUERIA ROSA", 18, 30),
                                 (26, "SUPERMERCADOS DIA", 20, 50)):
            imp = -round(random.uniform(lo, hi), 2)
            suma += imp
            tarjeta.append((mes(m, dia), com, imp))
        if m < 9:
            cuenta.append((mes(m + 1, 2), "LIQUIDACION TARJETA CREDITO 4321", round(suma, 2)))
    html_xls(d / "Movimientos_cuenta_20260930.xls",
             [(dmy(a), c, es(i), es(s)) for a, c, i, s in con_saldo(cuenta, 2310.45)],
             ["Fecha", "Concepto", "Importe", "Saldo"])
    xml_xls(d / "Tarjeta_4321_jul-sep.xls", [(dmy(a), c, es(i)) for a, c, i in tarjeta],
            ["Fecha operación", "Comercio", "Importe de la operación"])


def javier(d):
    """Programador. Cuenta en CSV cp1252 con columnas de más, y dos descargas
    que se solapan; tarjeta en .xlsx de verdad."""
    cuenta = []
    for m in (7, 8, 9):
        cuenta += [(mes(m, 28), "TRANSFERENCIA NOMINA TECHSOFT SL", 2380.00),
                   (mes(m, 1), "HIPOTECA PRESTAMO 0012", -710.35),
                   (mes(m, 6), "RECIBO DIGI SPAIN", -25.00),
                   (mes(m, 9), "RECIBO GIMNASIO BASIC FIT", -29.99),
                   (mes(m, 12), "BIZUM ENVIADO A CARLOS CENA", -23.50),
                   (mes(m, 15), "BIZUM RECIBIDO DE LAURA ENTRADAS", 40.00),
                   (mes(m, 18), "COMPRA TIENDANIMAL", -round(random.uniform(30, 60), 2)),
                   (mes(m, 20), "CLINICA VETERINARIA ARCA", -round(random.uniform(35, 90), 2)),
                   (mes(m, 25), "NETFLIX.COM SUSCRIPCION", -13.99)]
    cuenta.append((mes(8, 11), "COMPRA UDEMY CURSO PYTHON", -14.99))
    filas = con_saldo(cuenta, 5120.00)
    cab = ["F. Operación", "F. Valor", "Concepto", "Importe", "Saldo disponible"]
    fila = lambda a, c, i, s: (dmy(a), dmy(a + timedelta(days=1)), c, es(i), es(s))
    csv(d / "extracto_jul_ago.csv", [fila(*f) for f in filas if f[0].month in (7, 8)], cab)
    csv(d / "extracto_ago_sep.csv", [fila(*f) for f in filas if f[0].month in (8, 9)], cab)
    tarjeta = []
    for m in (7, 8, 9):
        for dia, com, lo, hi in ((3, "AMAZON MARKETPLACE", 15, 70), (13, "GLOVO PEDIDO", 12, 30),
                                 (19, "STEAM PURCHASE", 5, 40), (27, "REPSOL ESTACION", 40, 65)):
            tarjeta.append((mes(m, dia), com, -round(random.uniform(lo, hi), 2)))
    tarjeta.append((mes(8, 5), "DEVOLUCION AMAZON MARKETPLACE", 19.99))
    xlsx(d / "tarjeta_visa_2026.xlsx", [(dmy(a), c, i) for a, c, i in tarjeta],
         ["Fecha", "Comercio", "Importe de la operación"])


def marta(d):
    """Autónoma. Dos cuentas del mismo banco (negocio y personal) con
    traspasos entre ellas y muchos Bizum."""
    negocio, personal = [], []
    for m in (7, 8, 9):
        for dia, cli, imp in ((4, "ESTUDIO NORTE SL", 1210.00), (19, "LOPEZ Y ASOCIADOS", 847.00)):
            negocio.append((mes(m, dia), f"TRANSFERENCIA DE {cli}", imp + random.choice((0, 121.0))))
        negocio += [(mes(m, 30 if m != 9 else 29), "SEG SOCIAL REGIMEN AUTONOMOS", -294.00),
                    (mes(m, 10), "ADOBE SYSTEMS SOFTWARE", -60.49),
                    (mes(m, 21), "TRASPASO A CUENTA PERSONAL", -1200.00)]
        personal += [(mes(m, 21), "TRASPASO DESDE CUENTA NEGOCIO", 1200.00),
                     (mes(m, 2), "RECIBO ALQUILER PISO", -780.00),
                     (mes(m, 7), "COMPRA LIDL", -round(random.uniform(40, 90), 2)),
                     (mes(m, 16), "RECIBO ENDESA ENERGIA", -round(random.uniform(40, 70), 2))]
        for dia in (3, 9, 13, 24):
            personal.append((mes(m, dia), random.choice(("BIZUM DE ELENA CENA", "BIZUM DE PABLO REGALO")),
                             round(random.uniform(10, 30), 2)))
            personal.append((mes(m, dia + 1), random.choice(("BIZUM A ELENA CINE", "BIZUM A SARA BAR")),
                             -round(random.uniform(8, 25), 2)))
    negocio.append((mes(7, 20), "AEAT MODELO 303 IVA 2T", -612.30))
    cab = ["Fecha", "Descripción", "Importe (€)", "Saldo (€)"]
    html_xls(d / "cuenta_negocio_3T2026.xls",
             [(dmy(a), c, es(i), es(s)) for a, c, i, s in con_saldo(negocio, 3400.00)], cab)
    html_xls(d / "cuenta_personal_3T2026.xls",
             [(dmy(a), c, es(i), es(s)) for a, c, i, s in con_saldo(personal, 1150.00)], cab)


def pedro(d):
    """Jubilado. CSV sin columna de saldo, pensión que entra el día 1-2
    (es del mes anterior), recibos fijos, y dos descargas solapadas."""
    movs = []
    for m in (7, 8, 9, 10):
        movs.append((mes(m, 1 if m % 2 else 2), "ABONO PENSION INSS", 1432.18))
    for m in (7, 8, 9):
        movs += [(mes(m, 5), "RECIBO COMUNIDAD PROPIETARIOS", -65.00),
                 (mes(m, 6), "NATURGY IBERIA ADEUDO", -round(random.uniform(38, 60), 2)),
                 (mes(m, 9), "RECIBO TELEFONICA", -39.90),
                 (mes(m, 12), "RECIBO SANITAS SEGURO SALUD", -58.40),
                 (mes(m, 15), "COMPRA EROSKI", -round(random.uniform(80, 140), 2)),
                 (mes(m, 23), "FARMACIA LDO GARCIA", -round(random.uniform(10, 30), 2)),
                 (mes(m, 27), "TRANSFERENCIA A HIJA MARIA", -150.00)]
    movs.sort()
    cab = ["Fecha", "Fecha valor", "Concepto", "Importe"]
    fila = lambda a, c, i: (dmy(a), dmy(a), c, es(i))
    csv(d / "movimientos (1).csv", [fila(*f) for f in movs if f[0] < date(2026, 9, 1)], cab)
    csv(d / "movimientos (2).csv", [fila(*f) for f in movs if f[0] >= date(2026, 8, 1)], cab)


def ana(d):
    """Joven, solo tarjetas: una de un neobanco que exporta en inglés con
    punto decimal, otra española. Y dos ficheros que no son extractos."""
    rev = []
    for m in (7, 8, 9):
        for dia, com, lo, hi in ((2, "Spotify", 10.99, 10.99), (8, "Uber Eats", 12, 28),
                                 (14, "Primark", 15, 45), (20, "Ryanair", 30, 120),
                                 (26, "Mercadona", 25, 60)):
            rev.append((mes(m, dia), com, -round(random.uniform(lo, hi), 2)))
        rev.append((mes(m, 1), "Top-Up by *1234", 400.00))
    csv(d / "account-statement_2026-07-01_2026-09-30_es-es_a1b2c3.csv",
        [(f"{a.isoformat()} 10:15:00", f"{a.isoformat()} 10:16:00", c, f"{i:.2f}", "EUR", "COMPLETED")
         for a, c, i in sorted(rev)],
        ["Started Date", "Completed Date", "Description", "Amount", "Currency", "State"],
        sep=",", encoding="utf-8", preambulo=())
    tarj = [(mes(m, dia), com, -round(random.uniform(10, 70), 2))
            for m in (7, 8, 9) for dia, com in ((6, "EL CORTE INGLES"), (17, "DECATHLON"), (28, "FNAC"))]
    xml_xls(d / "tarjeta_debito.xls", [(dmy(a), c, es(i)) for a, c, i in tarj],
            ["Fecha operación", "Comercio", "Importe de la operación"])
    (d / "Recibo_alquiler_julio.pdf").write_bytes(b"%PDF-1.4\n% recibo de ejemplo\n")
    (d / "extracto_vacio.xls").write_bytes(b"")


# ---------------------------------------------------------------- ronda 2
# Datos más parecidos a un extracto real: el súper varias veces al mes con
# importes distintos, recibos que bailan unos días, devoluciones y algún
# cargo idéntico repetido de verdad (la deduplicación no debe comérselo).
def frecuentes(m, comercio, veces, lo, hi, primero=1, ultimo=28):
    dias = sorted(random.sample(range(primero, ultimo + 1), veces))
    return [(mes(m, d), comercio, -round(random.uniform(lo, hi), 2)) for d in dias]


def recibo(m, concepto, dia, importe, baile=3):
    """Un recibo que el banco pasa cerca de su día, no siempre el mismo."""
    return (mes(m, min(28, max(1, dia + random.randint(-baile, baile)))), concepto, importe)


def pareja(d):
    """Pareja con cuenta conjunta (CSV) y una tarjeta cada uno, del mismo
    banco y con el mismo formato: dos ficheros de tarjeta casi iguales y dos
    cafés idénticos el mismo día en cada una."""
    conj, t1, t2 = [], [], []
    for m in (7, 8, 9):
        conj += [recibo(m, "TRANSFERENCIA NOMINA INDRA SISTEMAS", 27, 1890.00, 1),
                 recibo(m, "TRANSFERENCIA NOMINA HOSPITAL QUIRON", 28, 1620.00, 1),
                 recibo(m, "ADEUDO PRESTAMO HIPOTECARIO", 1, -845.20, 0),
                 recibo(m, "RECIBO IBERDROLA CLIENTES", 6, -round(random.uniform(55, 95), 2)),
                 recibo(m, "RECIBO ESCUELA INFANTIL LOS PINOS", 3, -310.00),
                 recibo(m, "RECIBO MAPFRE SEGURO HOGAR", 15, -32.40),
                 recibo(m, "RECIBO ORANGE ESPAGNE", 20, -54.95)]
        conj += frecuentes(m, "COMPRA MERCADONA", 5, 35, 140)
        conj += frecuentes(m, "COMPRA ALCAMPO", 2, 60, 180)
        for tarjeta, final in ((t1, "1111"), (t2, "2222")):
            tarjeta += frecuentes(m, "CAFETERIA LA ESQUINA", 6, 1.8, 4.5)
            tarjeta += frecuentes(m, "GASOLINERA CEPSA", 2, 45, 70)
            tarjeta += frecuentes(m, "AMAZON MARKETPLACE", 2, 9, 60)
        t1 += frecuentes(m, "DECATHLON", 1, 20, 90)
        t2 += frecuentes(m, "PRIMARK", 1, 15, 60)
    # Los dos cafés iguales, el mismo día, en cada tarjeta: gastos reales.
    for t in (t1, t2):
        t += [(mes(8, 14), "CAFETERIA LA ESQUINA", -2.60)] * 2
    t2.append((mes(8, 22), "DEVOLUCION PRIMARK", 24.99))
    # Cada tarjeta se liquida en la conjunta el día 2 del mes siguiente.
    for t, final in ((t1, "1111"), (t2, "2222")):
        for m in (7, 8):
            suma = round(sum(i for a, _, i in t if a.month == m), 2)
            conj.append((mes(m + 1, 2), f"LIQUIDACION TARJETA {final}", suma))
    cab = ["Fecha", "Fecha valor", "Concepto", "Importe", "Saldo"]
    csv(d / "movimientos_cuenta_conjunta.csv",
        [(dmy(a), dmy(a), c, es(i), es(s)) for a, c, i, s in con_saldo(conj, 4200.00)], cab)
    for t, final in ((t1, "1111"), (t2, "2222")):
        xml_xls(d / f"Tarjeta_{final}_movimientos.xls",
                [(dmy(a), c, es(i)) for a, c, i in sorted(t)],
                ["Fecha operación", "Comercio", "Importe de la operación"])


def migrante(d):
    """Viene de otra app de finanzas: trae un CSV exportado de ella, con sus
    propias categorías, y los extractos del banco (con columnas Cargo y
    Abono separadas, como exportan algunos bancos)."""
    movs = []
    for m in (4, 5, 6, 7, 8, 9):
        movs += [recibo(m, "NOMINA ACCENTURE", 28, 2150.00, 1),
                 recibo(m, "RECIBO ALQUILER C/ MAYOR 12", 1, -750.00, 1),
                 recibo(m, "ADEUDO NATURGY", 12, -round(random.uniform(35, 70), 2)),
                 recibo(m, "RECIBO VODAFONE", 8, -35.00),
                 recibo(m, "SPOTIFY P1234ABC", 4, -11.99, 0)]
        movs += frecuentes(m, "COMPRA TARJ LIDL", 4, 15, 70)
        movs += frecuentes(m, "COMPRA TARJ BURGER KING", 2, 8, 16)
        movs += frecuentes(m, "RENFE CERCANIAS", 3, 2.4, 2.4)
    movs.append((mes(6, 18), "DEVOLUCION ADEUDO NATURGY", 48.20))
    antes = sorted(x for x in movs if x[0].month <= 6)
    despues = sorted(x for x in movs if x[0].month >= 7)
    cats = {"NOMINA": "Salario", "ALQUILER": "Vivienda", "NATURGY": "Facturas", "VODAFONE": "Facturas",
            "SPOTIFY": "Ocio", "LIDL": "Supermercado", "BURGER": "Restaurantes", "RENFE": "Transporte"}
    cat = lambda c: next((v for k, v in cats.items() if k in c), "Otros")
    csv(d / "export_otra_app_2026-06-30.csv",
        [(a.isoformat(), c.title(), cat(c), f"{i:.2f}") for a, c, i in antes],
        ["date", "description", "category", "amount"], sep=",", encoding="utf-8", preambulo=())
    saldo = 1800.00
    filas = []
    for a, c, i in despues:
        saldo = round(saldo + i, 2)
        filas.append((dmy(a), c, es(-i) if i < 0 else "", es(i) if i > 0 else "", es(saldo)))
    xlsx(d / "Extracto_cuenta_jul-sep_2026.xlsx", filas, ["Fecha", "Concepto", "Cargo", "Abono", "Saldo"])


def estudiante(d):
    """Estudiante con un único neobanco (CSV en inglés, otro formato que el
    de la ronda 1). Importes pequeños, Bizum y la ayuda de sus padres."""
    movs = []
    for m in (7, 8, 9):
        movs.append(recibo(m, "Transferencia de PADRES GARCIA", 1, 300.00, 2))
        movs += frecuentes(m, "CAFE BAR UNIVERSIDAD", 8, 1.2, 3.5)
        movs += frecuentes(m, "Mercadona", 3, 8, 35)
        movs += frecuentes(m, "Glovo", 2, 9, 22)
        movs += [(mes(m, 11), "Netflix", -5.49), recibo(m, "Tuenti Movil", 5, -10.00, 1),
                 recibo(m, "Bizum a MARIO piso gastos", 3, -120.00, 1)]
        for _ in range(3):
            movs.append((mes(m, random.randint(1, 28)), "Bizum de LUCAS cena", round(random.uniform(5, 20), 2)))
    movs += [(mes(9, 9), "CAFE BAR UNIVERSIDAD", -1.30)] * 2
    movs.append((mes(9, 15), "Beca MEC MINISTERIO EDUCACION", 1500.00))
    movs.append((mes(8, 3), "Zalando", -39.95))
    movs.append((mes(8, 19), "Zalando Refund", 39.95))
    cab = ["Date", "Payee", "Account number", "Transaction type", "Payment reference",
           "Amount (EUR)", "Amount (Foreign Currency)", "Type Foreign Currency", "Exchange Rate"]
    csv(d / "n26-csv-transactions.csv",
        [(a.isoformat(), c, "", "Presentment" if i < 0 else "Income", "", f"{i:.2f}", "", "", "")
         for a, c, i in sorted(movs)], cab, sep=",", encoding="utf-8", preambulo=())


def autonomo(d):
    """Autónomo de verdad: facturas cobradas, IVA trimestral (303), IRPF
    fraccionado (130), cuota, gastos deducibles y la devolución de la renta.
    Prueba la categoría Impuestos de la 2.12.1."""
    movs = []
    for m in (7, 8, 9):
        for _ in range(random.randint(2, 4)):
            movs.append((mes(m, random.randint(1, 28)), random.choice(
                ("TRANSF. DE REFORMAS GIL SL FRA", "TRANSF. DE COMUNIDAD PROP. SOL FRA",
                 "TRANSF. DE PARTICULAR ANTONIO RUIZ")) + f" 2026-{random.randint(10, 99)}",
                round(random.uniform(300, 2400), 2)))
        movs += [recibo(m, "CUOTA AUTONOMOS TGSS", 28, -320.00, 0),
                 recibo(m, "RECIBO GESTORIA ASESORES DEL SUR", 5, -60.50, 1),
                 recibo(m, "RECIBO SEGURO RC MAPFRE", 10, -41.20, 2),
                 recibo(m, "RECIBO ALQUILER NAVE POLIGONO", 1, -450.00, 1)]
        movs += frecuentes(m, "COMPRA LEROY MERLIN", 3, 25, 260)
        movs += frecuentes(m, "BRICOMART", 2, 40, 400)
        movs += frecuentes(m, "GASOLINERA REPSOL", 3, 50, 85)
    movs += [(mes(7, 18), "AEAT MODELO 303 2T 2026", -1184.62),
             (mes(7, 18), "AEAT MODELO 130 2T 2026", -402.10),
             (mes(7, 3), "DEVOLUCION AEAT IRPF 2025", 612.00),
             (mes(8, 12), "IMPUESTO VEHICULOS AYTO", -98.40),
             (mes(9, 5), "DEVOLUCION BRICOMART TICKET", 62.30)]
    cab = ["Fecha", "Concepto", "Importe", "Saldo"]
    html_xls(d / "Movimientos_cuenta_empresa.xls",
             [(dmy(a), c, es(i), es(s)) for a, c, i, s in con_saldo(movs, 6200.00)], cab)


def sync(d):
    """Solo quiere que los movimientos entren en SU Excel de siempre (con su
    hoja de datos y otra de totales), no el Excel de la herramienta."""
    movs = []
    for m in (7, 8, 9):
        movs += [recibo(m, "NOMINA AYUNTAMIENTO", 30 if m != 9 else 28, 1740.00, 0),
                 recibo(m, "RECIBO HIPOTECA", 2, -612.00, 0),
                 recibo(m, "RECIBO ENDESA", 9, -round(random.uniform(40, 80), 2)),
                 recibo(m, "RECIBO JAZZTEL", 14, -38.00)]
        movs += frecuentes(m, "COMPRA CONSUM", 4, 20, 95)
        movs += frecuentes(m, "FARMACIA", 2, 5, 25)
    cab = ["Fecha", "Concepto", "Importe", "Saldo"]
    html_xls(d / "extracto_julio_sept.xls",
             [(dmy(a), c, es(i), es(s)) for a, c, i, s in con_saldo(movs, 900.00)], cab)
    # Su contabilidad de siempre: una hoja de datos con lo de junio metido a
    # mano y otra de totales con fórmulas que apuntan a ella.
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"
    ws.append(["Fecha", "Concepto", "Importe", "Categoría"])
    for a, c, i, cat in ((date(2026, 6, 2), "RECIBO HIPOTECA", -612.00, "Piso"),
                         (date(2026, 6, 9), "RECIBO ENDESA", -55.10, "Luz/Agua"),
                         (date(2026, 6, 20), "COMPRA CONSUM", -64.35, "Comida"),
                         (date(2026, 6, 30), "NOMINA AYUNTAMIENTO", 1740.00, "Ingresos")):
        ws.append([a, c, i, cat])
    t = wb.create_sheet("Totales")
    t.append(["Categoría", "Total"])
    for i, cat in enumerate(("Piso", "Luz/Agua", "Comida", "Ingresos"), start=2):
        t.append([cat, f'=SUMIF(Datos!D:D,A{i},Datos!C:C)'])
    wb.save(d / "mi_contabilidad.xlsx")


RONDAS = {
    1: {"p1_lucia": lucia, "p2_javier": javier, "p3_marta": marta,
        "p4_pedro": pedro, "p5_ana": ana},
    # Marta y Ana repiten con los mismos objetivos para confirmar lo arreglado.
    2: {"p1_pareja": pareja, "p2_migrante": migrante, "p3_estudiante": estudiante,
        "p4_autonomo": autonomo, "p5_sync": sync, "p6_marta": marta, "p7_ana": ana},
}


def main():
    import sys
    args = sys.argv[1:]
    ronda = max(RONDAS)
    if args[:1] == ["--ronda"]:
        ronda, args = int(args[1]), args[2:]
    if not args:
        sys.exit(__doc__)
    destino = Path(args[0]).resolve()
    zips = sorted(RAIZ.glob("eledger_v*.zip"), key=lambda p: p.stat().st_mtime)
    zip_release = Path(args[1]) if len(args) > 1 else (zips[-1] if zips else None)
    web = Path(args[2]) if len(args) > 2 else RAIZ.parent / "eledger-web"
    if not zip_release or not zip_release.exists():
        sys.exit("No encuentro el ZIP de la release: genéralo con app/exportar.py "
                 "o pásalo como segundo argumento.")
    if not web.is_dir():
        sys.exit(f"No encuentro la web en {web}: pásala como tercer argumento.")
    print(f"Ronda {ronda} · release: {zip_release.name} · web: {web}")
    for nombre, generar in RONDAS[ronda].items():
        base = destino / nombre
        if base.exists():
            shutil.rmtree(base)
        (base / "descargas").mkdir(parents=True)
        with zipfile.ZipFile(zip_release) as z:
            z.extractall(base / "eledger")
        for f in (base / "eledger").rglob("*"):
            if f.suffix in (".sh", ".command"):
                f.chmod(0o755)
        shutil.copytree(web, base / "web", ignore=shutil.ignore_patterns(".git", "CLAUDE.md"))
        generar(base / "descargas")
        print(nombre, sorted(p.name for p in (base / "descargas").iterdir()))


if __name__ == "__main__":
    main()
